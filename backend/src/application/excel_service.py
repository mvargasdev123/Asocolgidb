import io
import openpyxl
from datetime import datetime, date
from typing import Dict, Any, List, Optional
from sqlmodel import Session, select

from domain.models.persona import Persona
from domain.models.roles import DatosAsociado, DatosVoluntario
from domain.models.expediente import Expediente
from domain.models.comentario import Comentario
from domain.models.catalogos import (
    Nacionalidad, Ciudad, NivelEducativo, MotivoConsulta, 
    Derivacion, TecnicaAcogida, TipoDocumento
)

def _clean_val(val: Any) -> Optional[str]:
    if val is None:
        return None
    val_str = str(val).strip()
    if not val_str or val_str.lower() in ["none", "null", "nan"]:
        return None
    return val_str

def _format_date_export(d: Any) -> Optional[str]:
    if d is None:
        return None
    if isinstance(d, datetime):
        d = d.date()
    if isinstance(d, date):
        return d.strftime("%m/%d/%Y")
    return str(d)

def _parse_date(val: Any) -> Optional[date]:
    if val is None:
        return None
    if isinstance(val, date) and not isinstance(val, datetime):
        return val
    if isinstance(val, datetime):
        return val.date()
    val_str = str(val).strip()
    if not val_str or val_str.lower() in ["none", "null", "nan", "n/a"]:
        return None
    if " " in val_str:
        val_str = val_str.split(" ")[0]
    elif "T" in val_str:
        val_str = val_str.split("T")[0]

    # Intenta formatos con prioridad Mes/Día/Año (MM/DD/YYYY, MM-DD-YYYY)
    for fmt in ("%m/%d/%Y", "%m-%d-%Y", "%m/%d/%y", "%m-%d-%y", "%Y-%m-%d", "%d/%m/%Y", "%d-%m-%Y", "%Y/%m/%d"):
        try:
            return datetime.strptime(val_str, fmt).date()
        except ValueError:
            continue
    return None



def _get_or_create_catalog(session: Session, model_class, name_attr: str, val: Optional[str]) -> Optional[int]:
    if not val:
        return None
    val_clean = val.strip()
    statement = select(model_class).where(getattr(model_class, name_attr) == val_clean)
    obj = session.exec(statement).first()
    if not obj:
        obj = model_class(**{name_attr: val_clean})
        session.add(obj)
        session.commit()
        session.refresh(obj)
    return obj.id

def _obtener_comentarios_texto(session: Session, id_persona: int, tipos: List[str]) -> Optional[str]:
    statement = select(Comentario).where(
        Comentario.id_persona == id_persona,
        Comentario.tipo.in_(tipos)
    )
    comentarios = session.exec(statement).all()
    if not comentarios:
        return None
    textos = [c.texto for c in comentarios if c.texto and c.texto.strip()]
    return " ; ".join(textos) if textos else None

def _guardar_comentario_si_existe(session: Session, id_persona: int, texto: Optional[str], tipo: str):
    if not texto:
        return
    texto_clean = str(texto).strip()
    if not texto_clean or texto_clean.lower() in ["none", "null", "nan", "n/a"]:
        return
    existente = session.exec(
        select(Comentario).where(
            Comentario.id_persona == id_persona,
            Comentario.tipo == tipo,
            Comentario.texto == texto_clean
        )
    ).first()
    if not existente:
        nuevo = Comentario(
            id_persona=id_persona,
            texto=texto_clean,
            tipo=tipo
        )
        session.add(nuevo)


def generar_excel_memoria(session: Session, tipo_exportacion: str = "completa") -> bytes:
    wb = openpyxl.Workbook()
    wb.remove(wb.active)

    # 1. Pestaña BD (31 columnas exactas a Modelo basededatos Asocolgi2026.xlsx)
    if tipo_exportacion in ["completa", "asociados", "voluntarios"]:
        ws_bd = wb.create_sheet(title="BD")
        headers_bd = [
            "FECHA ATENCIÓN", "TIPO IDENTIFICACION", "NÚMERO IDENTIFICACION", "NOMBRE COMPLETO",
            "TELEFONO", "E-MAIL", "MOTIVO CONSULTA", "No. REGISTRO ASOCIADO",
            "No. REGISTRO EXPEDIENTE", "No. REGISTRO VOLUNTARIO", "DIRECCIÓN", "CIUDAD RESIDENCIA",
            "CODIGO POSTAL", "FECHA NACIMIENTO", "PAIS NACIMIENTO", "GENERO", "SITUACIÓN ADMINISTRATIVA",
            "UNIDAD FAMILIAR", "MADRE SOLTERA", "VIOLENCIA GENERO", "TIENE PADRON", "FECHA PADRON",
            "CONTACTO EMERGENCIA NOMBRE", "CONTACTO EMERGENCIA PARENTESCO", "CONTACTO EMERGENCIA TELEFONO",
            "AUTORIZA PROTECCIÓN DATOS", "AUTORIZA USO IMAGEN", "DERIVACION", "NIVEL EDUCATIVO",
            "TECNICA ACOGIDA", "COMENTARIOS"
        ]
        ws_bd.append(headers_bd)

        query = select(Persona)
        if tipo_exportacion == "asociados":
            query = query.join(DatosAsociado)
        elif tipo_exportacion == "voluntarios":
            query = query.join(DatosVoluntario)

        personas = session.exec(query).all()
        for p in personas:
            tipo_doc_nombre = session.get(TipoDocumento, p.id_tipo_documento).nombre if p.id_tipo_documento else None
            nac_nombre = session.get(Nacionalidad, p.id_nacionalidad).nombre if p.id_nacionalidad else None
            motivo_nombre = session.get(MotivoConsulta, p.id_motivo_consulta).nombre if p.id_motivo_consulta else None
            ciudad_nombre = session.get(Ciudad, p.id_ciudad).nombre if p.id_ciudad else None
            derivacion_nombre = session.get(Derivacion, p.id_derivacion).nombre if p.id_derivacion else None
            nivel_edu_nombre = session.get(NivelEducativo, p.id_nivel_educativo).nombre if p.id_nivel_educativo else None
            tecnica_nombre = session.get(TecnicaAcogida, p.id_tecnica_acogida).nombre if p.id_tecnica_acogida else None

            reg_aso = f"ASO26{p.id:04d}" if p.datos_asociado else None
            reg_vol = p.id if p.datos_voluntario else None
            first_exp = p.expedientes[0].numero_registro if p.expedientes else None

            comentarios_bd = _obtener_comentarios_texto(session, p.id, ["Persona", "General", "BD"])

            ws_bd.append([
                _format_date_export(p.fecha_atencion),
                tipo_doc_nombre,
                p.numero_identificacion,
                p.nombre_completo,
                p.telefono_principal,
                p.correo_electronico,
                motivo_nombre,
                reg_aso,
                first_exp,
                reg_vol,
                p.direccion_residencia,
                ciudad_nombre,
                p.codigo_postal,
                _format_date_export(p.fecha_nacimiento),
                nac_nombre,
                p.genero,
                p.situacion_admin,
                p.unidad_familiar,
                p.madre_soltera,
                p.violencia_genero,
                "Si" if p.tiene_padron else ("No" if p.tiene_padron is False else None),
                _format_date_export(p.fecha_padron),
                p.contacto_emergencia_nombre,
                p.contacto_emergencia_parentesco,
                p.contacto_emergencia_telefono,
                "Si" if p.autoriza_datos else "No",
                "Si" if p.autoriza_imagen else "No",
                derivacion_nombre,
                nivel_edu_nombre,
                tecnica_nombre,
                comentarios_bd
            ])

    # 2. Pestaña EXP (12 columnas exactas)
    if tipo_exportacion in ["completa", "expedientes"]:
        ws_exp = wb.create_sheet(title="EXP")
        headers_exp = [
            "No. REGISTRO EXPEDIENTE", "NUMERO DE EXPEDIENTE", "TIPO DE TRÁMITE",
            "FECHA DE PRESENTACIÓN", "REPRESENTANTE LEGAL", "ESTADO DE EXPEDIENTE",
            "CONSULTORIO JURIDICO", "APORTE SOCIAL", "SOLICITANTE DE TRAMITES EXTRANJERIA",
            "TIENE ANTECEDENTES TRADUCIDOS Y APOSTILLADOS", "FECHA DE RESOLUCIÓN", "COMENTARIOS EXP"
        ]
        ws_exp.append(headers_exp)

        expedientes = session.exec(select(Expediente)).all()
        for e in expedientes:
            comentarios_exp = _obtener_comentarios_texto(session, e.id_persona, ["Expediente"])
            ws_exp.append([
                e.numero_registro,
                e.numero_expediente_asignado,
                e.tipo_tramite,
                _format_date_export(e.fecha_presentacion),
                e.representante_legal,
                e.estado,
                e.consultorio_juridico,
                e.aporte_social,
                "Si" if e.solicitante_extranjeria else ("No" if e.solicitante_extranjeria is False else None),
                "Si" if e.antecedentes_traducidos_y_apostillados else ("No" if e.antecedentes_traducidos_y_apostillados is False else None),
                _format_date_export(e.fecha_resolucion),
                comentarios_exp
            ])

    # 3. Pestaña ASO (7 columnas exactas)
    if tipo_exportacion in ["completa", "asociados"]:
        ws_aso = wb.create_sheet(title="ASO")
        headers_aso = [
            "No. REGISTRO ASOCIADO", "ESTADO", "METODO PAGO", "PAGOS CONFIRMADOS",
            "AUTORIZA GRUPO WHATSAPP", "FECHA DE VINCULACION", "COMENTARIOS ASO"
        ]
        ws_aso.append(headers_aso)

        asociados = session.exec(select(DatosAsociado)).all()
        for a in asociados:
            comentarios_aso = _obtener_comentarios_texto(session, a.id_persona, ["Asociado"])
            reg_aso = f"ASO26{a.id_persona:04d}"
            ws_aso.append([
                reg_aso,
                a.estado_membresia or "ALTA",
                a.metodo_pago or "EFECTIVO",
                a.estado_pago or "SI",
                "SI" if a.autoriza_whatsapp else "NO",
                _format_date_export(a.fecha_vinculacion),
                comentarios_aso
            ])

    # 4. Pestaña VOL (8 columnas exactas)
    if tipo_exportacion in ["completa", "voluntarios"]:
        ws_vol = wb.create_sheet(title="VOL")
        headers_vol = [
            "No. REGISTRO VOLUNTARIO", "CARGO", "CARTA COMPROMISO", "FECHA ALTA",
            "FECHA BAJA", "TIPO VOLUNTARIADO", "CAMPO DE ACCION", "COMENTARIOS VOL"
        ]
        ws_vol.append(headers_vol)

        voluntarios = session.exec(select(DatosVoluntario)).all()
        for v in voluntarios:
            comentarios_vol = _obtener_comentarios_texto(session, v.id_persona, ["Voluntario"])
            carta_str = "ok-Firmado" if v.carta_compromiso_firmada else None
            ws_vol.append([
                v.id_persona,
                v.cargo,
                carta_str,
                _format_date_export(v.fecha_alta) or _format_date_export(v.fecha_vinculacion),
                _format_date_export(v.fecha_baja),
                v.tipo,
                v.campo_accion,
                comentarios_vol
            ])

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    return output.read()


def analizar_excel_dry_run(file_bytes: bytes, session: Session) -> Dict[str, Any]:
    output = {
        "total_filas": 0,
        "nuevos_listos_para_guardar": 0,
        "duplicados": [],
        "errores": []
    }

    try:
        excel_file = io.BytesIO(file_bytes)
        wb = openpyxl.load_workbook(excel_file, data_only=True)
    except Exception as e:
        output["errores"].append({"fila": 0, "mensaje": f"El archivo cargado no es un Excel válido: {str(e)}"})
        return output

    if "BD" not in wb.sheetnames:
        output["errores"].append({"fila": 0, "mensaje": "Falta la pestaña principal 'BD' en el archivo Excel."})
        return output

    ws_bd = wb["BD"]
    rows = list(ws_bd.iter_rows(values_only=True))
    if not rows or len(rows) < 2:
        return output

    headers = [str(h).strip().upper() if h else "" for h in rows[0]]
    try:
        idx_identificacion = headers.index("NÚMERO IDENTIFICACION")
    except ValueError:
        output["errores"].append({"fila": 1, "mensaje": "No se encontró la columna requerida 'NÚMERO IDENTIFICACION' en la pestaña BD."})
        return output

    idx_nombre = headers.index("NOMBRE COMPLETO") if "NOMBRE COMPLETO" in headers else -1

    for row_idx, row in enumerate(rows[1:], start=2):
        if not any(row):
            continue
        output["total_filas"] += 1

        num_ident = _clean_val(row[idx_identificacion]) if idx_identificacion < len(row) else None
        nombre = _clean_val(row[idx_nombre]) if idx_nombre != -1 and idx_nombre < len(row) else f"Fila {row_idx}"

        if not num_ident:
            output["errores"].append({
                "fila": row_idx,
                "mensaje": f"Fila {row_idx}: 'NÚMERO IDENTIFICACION' está vacío."
            })
            continue

        persona_existente = session.exec(
            select(Persona).where(Persona.numero_identificacion == num_ident)
        ).first()

        if persona_existente:
            output["duplicados"].append({
                "fila_excel": row_idx,
                "identificacion": num_ident,
                "nombre_excel": nombre or persona_existente.nombre_completo,
                "accion_recomendada": "resolver"
            })
        else:
            output["nuevos_listos_para_guardar"] += 1

    return output


def ejecutar_importacion_definitiva(file_bytes: bytes, decisiones: Dict[str, str], session: Session) -> Dict[str, Any]:
    excel_file = io.BytesIO(file_bytes)
    wb = openpyxl.load_workbook(excel_file, data_only=True)

    ws_bd = wb["BD"]
    bd_rows = list(ws_bd.iter_rows(values_only=True))
    bd_headers = [str(h).strip().upper() if h else "" for h in bd_rows[0]]

    def _get_bd_val(row, header_name):
        for h in bd_headers:
            if header_name.upper() in h:
                idx = bd_headers.index(h)
                if idx < len(row):
                    return _clean_val(row[idx])
        return None

    creados = 0
    actualizados = 0
    omitidos = 0

    identificacion_to_persona_id = {}
    bd_index_to_persona_id = {}
    reg_aso_to_persona_id = {}
    reg_exp_to_persona_id = {}
    reg_vol_to_persona_id = {}

    bd_aso_requested = set()
    bd_vol_requested = set()

    for idx, row in enumerate(bd_rows[1:], start=1):
        if not any(row):
            continue
        num_ident = _get_bd_val(row, "NÚMERO IDENTIFICACION") or _get_bd_val(row, "NUMERO IDENTIFICACION")
        if not num_ident:
            continue

        decision = decisiones.get(num_ident, "sobreescribir")
        if decision == "omitir":
            omitidos += 1
            p_ex = session.exec(select(Persona).where(Persona.numero_identificacion == num_ident)).first()
            if p_ex:
                identificacion_to_persona_id[num_ident] = p_ex.id
                bd_index_to_persona_id[idx] = p_ex.id
            continue

        nombre = _get_bd_val(row, "NOMBRE COMPLETO") or f"Persona {num_ident}"
        tipo_doc = _get_bd_val(row, "TIPO IDENTIFICACION")
        nac_str = _get_bd_val(row, "PAIS NACIMIENTO")
        ciudad_str = _get_bd_val(row, "CIUDAD RESIDENCIA")
        motivo_str = _get_bd_val(row, "MOTIVO CONSULTA")
        derivacion_str = _get_bd_val(row, "DERIVACION")
        nivel_edu_str = _get_bd_val(row, "NIVEL EDUCATIVO")
        tecnica_str = _get_bd_val(row, "TECNICA ACOGIDA")

        tipo_doc_id = _get_or_create_catalog(session, TipoDocumento, "nombre", tipo_doc) if tipo_doc else None
        nac_id = _get_or_create_catalog(session, Nacionalidad, "nombre", nac_str) if nac_str else None
        ciudad_id = _get_or_create_catalog(session, Ciudad, "nombre", ciudad_str) if ciudad_str else None
        motivo_id = _get_or_create_catalog(session, MotivoConsulta, "nombre", motivo_str) if motivo_str else None
        derivacion_id = _get_or_create_catalog(session, Derivacion, "nombre", derivacion_str) if derivacion_str else None
        nivel_edu_id = _get_or_create_catalog(session, NivelEducativo, "nombre", nivel_edu_str) if nivel_edu_str else None
        tecnica_id = _get_or_create_catalog(session, TecnicaAcogida, "nombre", tecnica_str) if tecnica_str else None

        f_atenc = _parse_date(_get_bd_val(row, "FECHA ATENCIÓN") or _get_bd_val(row, "FECHA ATENCION"))
        tel_princ = _get_bd_val(row, "TELEFONO") or _get_bd_val(row, "TELÉFONO")
        f_nac = _parse_date(_get_bd_val(row, "FECHA NACIMIENTO"))
        genero = _get_bd_val(row, "GENERO") or _get_bd_val(row, "GÉNERO")
        correo = _get_bd_val(row, "E-MAIL") or _get_bd_val(row, "EMAIL")
        direccion = _get_bd_val(row, "DIRECCIÓN") or _get_bd_val(row, "DIRECCION")
        cp = _get_bd_val(row, "CODIGO POSTAL")
        sit_admin = _get_bd_val(row, "SITUACIÓN ADMINISTRATIVA") or _get_bd_val(row, "SITUACION ADMINISTRATIVA")
        unid_fam_raw = _get_bd_val(row, "UNIDAD FAMILIAR")
        unid_fam = int(unid_fam_raw) if unid_fam_raw and str(unid_fam_raw).isdigit() else 1
        madre_solt = _get_bd_val(row, "MADRE SOLTERA")
        viol_gen = _get_bd_val(row, "VIOLENCIA GENERO")
        tiene_padron_val = _get_bd_val(row, "TIENE PADRON")
        tiene_padron = tiene_padron_val in ["Sí", "SI", "Si", "True", "true", "1"] if tiene_padron_val else None
        f_padron = _parse_date(_get_bd_val(row, "FECHA PADRON"))
        aut_datos = _get_bd_val(row, "AUTORIZA PROTECCIÓN DATOS") in ["Sí", "SI", "Si", "True", "true", "1"]
        aut_img = _get_bd_val(row, "AUTORIZA USO IMAGEN") in ["Sí", "SI", "Si", "True", "true", "1"]

        ce_nombre = _get_bd_val(row, "CONTACTO EMERGENCIA NOMBRE")
        ce_parent = _get_bd_val(row, "CONTACTO EMERGENCIA PARENTESCO")
        ce_tel = _get_bd_val(row, "CONTACTO EMERGENCIA TELEFONO")

        reg_aso_val = _get_bd_val(row, "REGISTRO ASOCIADO")
        reg_exp_val = _get_bd_val(row, "REGISTRO EXPEDIENTE")
        reg_vol_val = _get_bd_val(row, "REGISTRO VOLUNTARIO")

        comentario_bd = _get_bd_val(row, "COMENTARIOS")

        persona = session.exec(select(Persona).where(Persona.numero_identificacion == num_ident)).first()
        if persona:
            persona.nombre_completo = nombre
            persona.id_tipo_documento = tipo_doc_id or persona.id_tipo_documento
            persona.id_nacionalidad = nac_id or persona.id_nacionalidad
            persona.id_ciudad = ciudad_id or persona.id_ciudad
            persona.id_motivo_consulta = motivo_id or persona.id_motivo_consulta
            persona.id_derivacion = derivacion_id or persona.id_derivacion
            persona.id_nivel_educativo = nivel_edu_id or persona.id_nivel_educativo
            persona.id_tecnica_acogida = tecnica_id or persona.id_tecnica_acogida
            persona.fecha_atencion = f_atenc or persona.fecha_atencion
            persona.telefono_principal = tel_princ or persona.telefono_principal
            persona.fecha_nacimiento = f_nac or persona.fecha_nacimiento
            persona.genero = genero or persona.genero
            persona.correo_electronico = correo or persona.correo_electronico
            persona.direccion_residencia = direccion or persona.direccion_residencia
            persona.codigo_postal = cp or persona.codigo_postal
            persona.situacion_admin = sit_admin or persona.situacion_admin
            persona.unidad_familiar = unid_fam
            persona.madre_soltera = madre_solt or persona.madre_soltera
            persona.violencia_genero = viol_gen or persona.violencia_genero
            persona.tiene_padron = tiene_padron if tiene_padron is not None else persona.tiene_padron
            persona.fecha_padron = f_padron or persona.fecha_padron
            persona.autoriza_datos = aut_datos
            persona.autoriza_imagen = aut_img
            persona.contacto_emergencia_nombre = ce_nombre or persona.contacto_emergencia_nombre
            persona.contacto_emergencia_parentesco = ce_parent or persona.contacto_emergencia_parentesco
            persona.contacto_emergencia_telefono = ce_tel or persona.contacto_emergencia_telefono
            actualizados += 1
        else:
            persona = Persona(
                numero_identificacion=num_ident,
                nombre_completo=nombre,
                id_tipo_documento=tipo_doc_id,
                id_nacionalidad=nac_id,
                id_ciudad=ciudad_id,
                id_motivo_consulta=motivo_id,
                id_derivacion=derivacion_id,
                id_nivel_educativo=nivel_edu_id,
                id_tecnica_acogida=tecnica_id,
                fecha_atencion=f_atenc,
                telefono_principal=tel_princ,
                fecha_nacimiento=f_nac,
                genero=genero,
                correo_electronico=correo,
                direccion_residencia=direccion,
                codigo_postal=cp,
                situacion_admin=sit_admin,
                unidad_familiar=unid_fam,
                madre_soltera=madre_solt,
                violencia_genero=viol_gen,
                tiene_padron=tiene_padron,
                fecha_padron=f_padron,
                autoriza_datos=aut_datos,
                autoriza_imagen=aut_img,
                contacto_emergencia_nombre=ce_nombre,
                contacto_emergencia_parentesco=ce_parent,
                contacto_emergencia_telefono=ce_tel,
                activo=True
            )
            session.add(persona)
            session.commit()
            session.refresh(persona)
            creados += 1

        identificacion_to_persona_id[num_ident] = persona.id
        bd_index_to_persona_id[idx] = persona.id

        if reg_aso_val and str(reg_aso_val).upper() not in ["N/A", "NONE", "NO", "FALSE"]:
            reg_aso_to_persona_id[str(reg_aso_val).strip()] = persona.id
            bd_aso_requested.add(persona.id)
        if reg_exp_val and str(reg_exp_val).upper() not in ["N/A", "NONE", "NO", "FALSE"]:
            reg_exp_to_persona_id[str(reg_exp_val).strip()] = persona.id
        if reg_vol_val and str(reg_vol_val).upper() not in ["N/A", "NONE", "NO", "FALSE"]:
            reg_vol_to_persona_id[str(reg_vol_val).strip()] = persona.id
            bd_vol_requested.add(persona.id)

        # Guardar comentario etiquetado como "Persona" / "BD"
        _guardar_comentario_si_existe(session, persona.id, comentario_bd, "Persona")

    # 2. Pestaña ASO
    if "ASO" in wb.sheetnames:
        ws_aso = wb["ASO"]
        aso_rows = list(ws_aso.iter_rows(values_only=True))
        if len(aso_rows) > 1:
            aso_headers = [str(h).strip().upper() if h else "" for h in aso_rows[0]]

            def _get_aso_val(row, header_name):
                for h in aso_headers:
                    if header_name.upper() in h:
                        idx = aso_headers.index(h)
                        if idx < len(row):
                            return _clean_val(row[idx])
                return None

            for aso_idx, row in enumerate(aso_rows[1:], start=1):
                if not any(row): continue
                reg_aso = _get_aso_val(row, "REGISTRO ASOCIADO")
                num_ident_aso = _get_aso_val(row, "IDENTIFICACION")

                persona_id = None
                if num_ident_aso and num_ident_aso in identificacion_to_persona_id:
                    persona_id = identificacion_to_persona_id[num_ident_aso]
                elif reg_aso and str(reg_aso).strip() in reg_aso_to_persona_id:
                    persona_id = reg_aso_to_persona_id[str(reg_aso).strip()]
                elif reg_aso and str(reg_aso).startswith("ASO-"):
                    raw_id = str(reg_aso).split("-")[1]
                    if raw_id.isdigit() and int(raw_id) in bd_index_to_persona_id.values():
                        persona_id = int(raw_id)
                    elif raw_id.isdigit() and int(raw_id) in bd_index_to_persona_id:
                        persona_id = bd_index_to_persona_id[int(raw_id)]
                if not persona_id and aso_idx in bd_index_to_persona_id:
                    persona_id = bd_index_to_persona_id[aso_idx]

                if persona_id and session.get(Persona, persona_id):
                    aso_obj = session.exec(select(DatosAsociado).where(DatosAsociado.id_persona == persona_id)).first()
                    est_mem = _get_aso_val(row, "ESTADO") or "ALTA"
                    met_pag = _get_aso_val(row, "METODO PAGO") or "EFECTIVO"
                    est_pag = _get_aso_val(row, "PAGOS CONFIRMADOS") or "SI"
                    aut_wh_val = _get_aso_val(row, "AUTORIZA GRUPO WHATSAPP")
                    aut_wh = aut_wh_val in ["Sí", "SI", "Si", "True", "true", "1"] if aut_wh_val else True
                    f_vinc = _parse_date(_get_aso_val(row, "FECHA DE VINCULACION") or _get_aso_val(row, "VINCULACION"))
                    comentario_aso = _get_aso_val(row, "COMENTARIOS")

                    if not aso_obj:
                        aso_obj = DatosAsociado(
                            id_persona=persona_id,
                            estado_membresia=est_mem,
                            metodo_pago=met_pag,
                            estado_pago=est_pag,
                            autoriza_whatsapp=aut_wh,
                            fecha_vinculacion=f_vinc
                        )
                        session.add(aso_obj)
                    else:
                        aso_obj.estado_membresia = est_mem or aso_obj.estado_membresia
                        aso_obj.metodo_pago = met_pag or aso_obj.metodo_pago
                        aso_obj.estado_pago = est_pag or aso_obj.estado_pago
                        aso_obj.autoriza_whatsapp = aut_wh
                        if f_vinc:
                            aso_obj.fecha_vinculacion = f_vinc

                    _guardar_comentario_si_existe(session, persona_id, comentario_aso, "Asociado")

    # 3. Pestaña VOL
    if "VOL" in wb.sheetnames:
        ws_vol = wb["VOL"]
        vol_rows = list(ws_vol.iter_rows(values_only=True))
        if len(vol_rows) > 1:
            vol_headers = [str(h).strip().upper() if h else "" for h in vol_rows[0]]

            def _get_vol_val(row, header_name):
                for h in vol_headers:
                    if header_name.upper() in h:
                        idx = vol_headers.index(h)
                        if idx < len(row):
                            return _clean_val(row[idx])
                return None

            for vol_idx, row in enumerate(vol_rows[1:], start=1):
                if not any(row): continue
                reg_vol = _get_vol_val(row, "REGISTRO VOLUNTARIO")
                num_ident_vol = _get_vol_val(row, "IDENTIFICACION")

                persona_id = None
                if num_ident_vol and num_ident_vol in identificacion_to_persona_id:
                    persona_id = identificacion_to_persona_id[num_ident_vol]
                elif reg_vol and str(reg_vol).strip() in reg_vol_to_persona_id:
                    persona_id = reg_vol_to_persona_id[str(reg_vol).strip()]
                elif reg_vol and str(reg_vol).isdigit() and int(reg_vol) in bd_index_to_persona_id.values():
                    persona_id = int(reg_vol)
                elif reg_vol and str(reg_vol).isdigit() and int(reg_vol) in bd_index_to_persona_id:
                    persona_id = bd_index_to_persona_id[int(reg_vol)]
                if not persona_id and vol_idx in bd_index_to_persona_id:
                    persona_id = bd_index_to_persona_id[vol_idx]

                if persona_id and session.get(Persona, persona_id):
                    vol_obj = session.exec(select(DatosVoluntario).where(DatosVoluntario.id_persona == persona_id)).first()
                    cargo = _get_vol_val(row, "CARGO") or "Voluntario"
                    campo = _get_vol_val(row, "CAMPO DE ACCION") or "General"
                    tipo = _get_vol_val(row, "TIPO VOLUNTARIADO") or "Presencial"
                    f_alta = _parse_date(_get_vol_val(row, "FECHA ALTA"))
                    f_baja = _parse_date(_get_vol_val(row, "FECHA BAJA"))
                    carta_raw = _get_vol_val(row, "CARTA COMPROMISO")
                    carta = carta_raw is not None and carta_raw.upper() not in ["NO", "FALSE", "0", "NONE"]
                    comentario_vol = _get_vol_val(row, "COMENTARIOS")

                    if not vol_obj:
                        vol_obj = DatosVoluntario(
                            id_persona=persona_id,
                            cargo=cargo,
                            campo_accion=campo,
                            tipo=tipo,
                            horas_semana=1,
                            fecha_alta=f_alta,
                            fecha_baja=f_baja,
                            carta_compromiso_firmada=carta,
                            formulario_inscripcion=False
                        )
                        session.add(vol_obj)
                    else:
                        vol_obj.cargo = cargo or vol_obj.cargo
                        vol_obj.campo_accion = campo or vol_obj.campo_accion
                        vol_obj.tipo = tipo or vol_obj.tipo
                        if f_alta: vol_obj.fecha_alta = f_alta
                        if f_baja: vol_obj.fecha_baja = f_baja
                        vol_obj.carta_compromiso_firmada = carta

                    _guardar_comentario_si_existe(session, persona_id, comentario_vol, "Voluntario")

    # Crear registros salvaguarda para asociados y voluntarios indicados en BD
    for pid in bd_aso_requested:
        if session.get(Persona, pid):
            aso_ex = session.exec(select(DatosAsociado).where(DatosAsociado.id_persona == pid)).first()
            if not aso_ex:
                session.add(DatosAsociado(
                    id_persona=pid,
                    estado_membresia="ALTA",
                    metodo_pago="EFECTIVO",
                    estado_pago="SI",
                    autoriza_whatsapp=True
                ))

    for pid in bd_vol_requested:
        if session.get(Persona, pid):
            vol_ex = session.exec(select(DatosVoluntario).where(DatosVoluntario.id_persona == pid)).first()
            if not vol_ex:
                session.add(DatosVoluntario(
                    id_persona=pid,
                    cargo="Voluntario",
                    campo_accion="General",
                    tipo="Presencial",
                    horas_semana=1,
                    carta_compromiso_firmada=False,
                    formulario_inscripcion=False
                ))

    # 4. Pestaña EXP
    if "EXP" in wb.sheetnames:
        ws_exp = wb["EXP"]
        exp_rows = list(ws_exp.iter_rows(values_only=True))
        if len(exp_rows) > 1:
            exp_headers = [str(h).strip().upper() if h else "" for h in exp_rows[0]]

            def _get_exp_val(row, header_name):
                for h in exp_headers:
                    if header_name.upper() in h:
                        idx = exp_headers.index(h)
                        if idx < len(row):
                            return _clean_val(row[idx])
                return None

            for exp_idx, row in enumerate(exp_rows[1:], start=1):
                if not any(row): continue
                reg_exp = _get_exp_val(row, "REGISTRO EXPEDIENTE")
                num_ident_exp = _get_exp_val(row, "IDENTIFICACION")
                if not reg_exp and not num_ident_exp:
                    reg_exp = str(exp_idx)

                persona_id = None
                if num_ident_exp and num_ident_exp in identificacion_to_persona_id:
                    persona_id = identificacion_to_persona_id[num_ident_exp]
                elif reg_exp and str(reg_exp).strip() in reg_exp_to_persona_id:
                    persona_id = reg_exp_to_persona_id[str(reg_exp).strip()]
                elif reg_exp and str(reg_exp).isdigit() and int(reg_exp) in bd_index_to_persona_id.values():
                    persona_id = int(reg_exp)
                elif reg_exp and str(reg_exp).isdigit() and int(reg_exp) in bd_index_to_persona_id:
                    persona_id = bd_index_to_persona_id[int(reg_exp)]
                if not persona_id and exp_idx in bd_index_to_persona_id:
                    persona_id = bd_index_to_persona_id[exp_idx]

                if not persona_id:
                    # Usar la primera persona registrada como fallback
                    first_p = session.exec(select(Persona)).first()
                    if first_p:
                        persona_id = first_p.id

                if persona_id and session.get(Persona, persona_id):
                    reg_exp_str = str(reg_exp) if reg_exp else f"EXP-{exp_idx}"
                    exp_obj = session.exec(select(Expediente).where(Expediente.numero_registro == reg_exp_str)).first()
                    
                    tipo_tr = _get_exp_val(row, "TIPO DE TRÁMITE") or _get_exp_val(row, "TIPO TRÁMITE") or "Trámite General"
                    f_pres = _parse_date(_get_exp_val(row, "FECHA DE PRESENTACIÓN") or _get_exp_val(row, "PRESENTACIÓN")) or date.today()
                    est_exp = _get_exp_val(row, "ESTADO DE EXPEDIENTE") or _get_exp_val(row, "ESTADO") or "Favorable"
                    exp_asig = _get_exp_val(row, "NUMERO DE EXPEDIENTE") or _get_exp_val(row, "EXPEDIENTE ASIGNADO")
                    rep_leg = _get_exp_val(row, "REPRESENTANTE LEGAL")
                    cons_jur = _get_exp_val(row, "CONSULTORIO JURIDICO")
                    ap_soc = _get_exp_val(row, "APORTE SOCIAL") or "Sí"
                    sol_ext_val = _get_exp_val(row, "SOLICITANTE DE TRAMITES EXTRANJERIA")
                    sol_ext = sol_ext_val in ["Sí", "SI", "Si", "True", "true", "1"] if sol_ext_val else False
                    ant_ap_val = _get_exp_val(row, "TIENE ANTECEDENTES TRADUCIDOS Y APOSTILLADOS")
                    ant_ap = ant_ap_val in ["Sí", "SI", "Si", "True", "true", "1"] if ant_ap_val else False
                    f_res = _parse_date(_get_exp_val(row, "FECHA DE RESOLUCIÓN") or _get_exp_val(row, "RESOLUCIÓN"))
                    comentario_exp = _get_exp_val(row, "COMENTARIOS")

                    if not exp_obj:
                        exp_obj = Expediente(
                            id_persona=persona_id,
                            numero_registro=reg_exp_str,
                            tipo_tramite=tipo_tr,
                            fecha_presentacion=f_pres,
                            estado=est_exp,
                            numero_expediente_asignado=exp_asig,
                            representante_legal=rep_leg,
                            consultorio_juridico=cons_jur,
                            aporte_social=ap_soc,
                            solicitante_extranjeria=sol_ext,
                            antecedentes_traducidos_y_apostillados=ant_ap,
                            fecha_resolucion=f_res
                        )
                        session.add(exp_obj)
                    else:
                        exp_obj.tipo_tramite = tipo_tr or exp_obj.tipo_tramite
                        exp_obj.fecha_presentacion = f_pres or exp_obj.fecha_presentacion
                        exp_obj.estado = est_exp or exp_obj.estado
                        exp_obj.numero_expediente_asignado = exp_asig or exp_obj.numero_expediente_asignado
                        exp_obj.representante_legal = rep_leg or exp_obj.representante_legal
                        exp_obj.consultorio_juridico = cons_jur or exp_obj.consultorio_juridico
                        exp_obj.aporte_social = ap_soc or exp_obj.aporte_social
                        exp_obj.solicitante_extranjeria = sol_ext
                        exp_obj.antecedentes_traducidos_y_apostillados = ant_ap
                        if f_res:
                            exp_obj.fecha_resolucion = f_res

                    _guardar_comentario_si_existe(session, persona_id, comentario_exp, "Expediente")

    session.commit()

    return {
        "status": "success",
        "mensaje": f"Importación completada con éxito. Creados: {creados}, Actualizados: {actualizados}, Omitidos: {omitidos}.",
        "registros_creados": creados,
        "registros_actualizados": actualizados,
        "registros_omitidos": omitidos
    }
