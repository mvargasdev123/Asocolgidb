import io
import openpyxl
import pytest
from fastapi.testclient import TestClient
from sqlmodel import Session

from domain.models.persona import Persona
from application.excel_service import generar_excel_memoria
from application.auth_service import AuthService
from infrastructure.auth_repository import AuthRepository

def test_excel_export_email_sin_smtp_falla(client: TestClient, session: Session, usuario_prueba):
    auth_service = AuthService(AuthRepository(session))
    token = auth_service.crear_token_acceso({"sub": str(usuario_prueba.id)})
    headers = {"Authorization": f"Bearer {token}"}
    response = client.post("/excel/exportar", json={"tipo_exportacion": "completa"}, headers=headers)
    assert response.status_code == 400
    assert "SMTP" in response.json()["detail"]


def test_excel_descargar_directo(client: TestClient, session: Session, usuario_prueba):
    auth_service = AuthService(AuthRepository(session))
    token = auth_service.crear_token_acceso({"sub": str(usuario_prueba.id)})
    persona = Persona(
        numero_identificacion="Y1234567A",
        nombre_completo="Test Export Persona",
        correo_electronico="export@test.com",
        activo=True
    )
    session.add(persona)
    session.commit()

    headers = {"Authorization": f"Bearer {token}"}
    response = client.post("/excel/descargar-directo", json={"tipo_exportacion": "completa"}, headers=headers)
    assert response.status_code == 200
    assert response.headers["content-type"] == "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    assert len(response.content) > 0

    wb = openpyxl.load_workbook(io.BytesIO(response.content))
    assert "BD" in wb.sheetnames
    assert "ASO" in wb.sheetnames
    assert "EXP" in wb.sheetnames
    assert "VOL" in wb.sheetnames

    # Verificar encabezados exactos de Modelo basededatos Asocolgi2026.xlsx en fila 1
    headers_bd = [cell.value for cell in wb["BD"][1]]
    assert "DERIVACION" in headers_bd
    assert "COMENTARIOS" in headers_bd

    headers_aso = [cell.value for cell in wb["ASO"][1]]
    assert "No. REGISTRO ASOCIADO" in headers_aso
    assert "FECHA DE VINCULACION" in headers_aso
    assert "COMENTARIOS ASO" in headers_aso

    headers_exp = [cell.value for cell in wb["EXP"][1]]
    assert "No. REGISTRO EXPEDIENTE" in headers_exp
    assert "COMENTARIOS EXP" in headers_exp

    headers_vol = [cell.value for cell in wb["VOL"][1]]
    assert "No. REGISTRO VOLUNTARIO" in headers_vol
    assert "COMENTARIOS VOL" in headers_vol


def test_excel_export_subsets(session: Session):
    bytes_aso = generar_excel_memoria(session, tipo_exportacion="asociados")
    wb_aso = openpyxl.load_workbook(io.BytesIO(bytes_aso))
    assert "BD" in wb_aso.sheetnames
    assert "ASO" in wb_aso.sheetnames
    assert "EXP" not in wb_aso.sheetnames

    bytes_exp = generar_excel_memoria(session, tipo_exportacion="expedientes")
    wb_exp = openpyxl.load_workbook(io.BytesIO(bytes_exp))
    assert "EXP" in wb_exp.sheetnames
    assert "ASO" not in wb_exp.sheetnames


def test_export_mm_dd_yyyy_format(session: Session):
    from datetime import date
    p = Persona(
        numero_identificacion="Z99988877A",
        nombre_completo="Export Date Test",
        fecha_nacimiento=date(1955, 11, 13),
        activo=True
    )
    session.add(p)
    session.commit()

    bytes_out = generar_excel_memoria(session, tipo_exportacion="completa")
    wb = openpyxl.load_workbook(io.BytesIO(bytes_out))
    ws = wb["BD"]
    
    found_date = None
    for row in ws.iter_rows(values_only=True):
        if len(row) > 3 and row[2] == "Z99988877A":
            found_date = row[13] # FECHA NACIMIENTO column index 13
            break
    assert found_date == "11/13/1955"

